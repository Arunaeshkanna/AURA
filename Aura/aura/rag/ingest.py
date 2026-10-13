from pathlib import Path

import requests
from bs4 import BeautifulSoup

from aura.core.config import get_settings
from aura.rag.chunking import chunk_text
from aura.rag.store import HybridStore


class Ingestor:
    def __init__(self) -> None:
        self.settings = get_settings()
        self.store = HybridStore()

    def ingest_text(self, title: str, content: str, source: str = "manual", metadata: dict | None = None) -> int:
        chunks = chunk_text(title, content, source, metadata or {})
        return self.store.add(chunks)

    def ingest_upload(self, filename: str, data: bytes) -> int:
        text = self.extract_file_text(filename, data)
        if not text.strip():
            return 0
        return self.ingest_text(
            Path(filename).stem,
            text,
            "upload",
            {
                "filename": filename,
                "extension": Path(filename).suffix.lower(),
                "collection": "uploads",
                "uploaded_kind": "document",
            },
        )

    def ingest_url(self, url: str, title: str | None = None, metadata: dict | None = None) -> int:
        response = requests.get(
            url,
            timeout=20,
            headers={"User-Agent": "AURA research agent"},
        )
        response.raise_for_status()
        soup = BeautifulSoup(response.text, "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        page_title = title or (soup.title.string.strip() if soup.title and soup.title.string else url)
        text = "\n".join(line.strip() for line in soup.get_text("\n").splitlines() if line.strip())
        return self.ingest_text(
            page_title,
            text,
            "url",
            {
                **(metadata or {}),
                "url": url,
                "collection": "web_uploads",
                "uploaded_kind": "url",
            },
        )

    def extract_file_text(self, filename: str, data: bytes) -> str:
        suffix = Path(filename).suffix.lower()
        if suffix in {".txt", ".md", ".csv", ".json", ".html", ".htm"}:
            text = data.decode("utf-8", errors="ignore")
            if suffix in {".html", ".htm"}:
                return BeautifulSoup(text, "html.parser").get_text("\n")
            return text
        if suffix == ".pdf":
            return self._extract_pdf(data)
        if suffix in {".docx", ".doc"}:
            return self._extract_docx(data)
        if suffix in {".pptx", ".ppt"}:
            return self._extract_pptx(data)
        if suffix in {".xlsx", ".xls"}:
            return self._extract_excel(data)
        return data.decode("utf-8", errors="ignore")

    def _extract_pdf(self, data: bytes) -> str:
        from io import BytesIO

        try:
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            return f"PDF extraction failed: {exc}"

    def _extract_docx(self, data: bytes) -> str:
        from io import BytesIO

        try:
            from docx import Document

            document = Document(BytesIO(data))
            return "\n".join(paragraph.text for paragraph in document.paragraphs)
        except Exception as exc:
            return f"Word extraction failed: {exc}"

    def _extract_pptx(self, data: bytes) -> str:
        from io import BytesIO

        try:
            from pptx import Presentation

            presentation = Presentation(BytesIO(data))
            lines = []
            for slide_number, slide in enumerate(presentation.slides, start=1):
                lines.append(f"Slide {slide_number}")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        lines.append(shape.text)
            return "\n".join(lines)
        except Exception as exc:
            return f"PowerPoint extraction failed: {exc}"

    def _extract_excel(self, data: bytes) -> str:
        from io import BytesIO

        try:
            import pandas as pd

            sheets = pd.read_excel(BytesIO(data), sheet_name=None)
            blocks = []
            for name, frame in sheets.items():
                blocks.append(f"Sheet: {name}")
                blocks.append(frame.head(100).to_csv(index=False))
            return "\n".join(blocks)
        except Exception as exc:
            return f"Excel extraction failed: {exc}"

    def ingest_directory(self, directory: Path | None = None) -> int:
        root = directory or self.settings.docs_dir
        total = 0
        for path in root.glob("**/*"):
            if path.is_file() and path.suffix.lower() in {".txt", ".md", ".csv"}:
                content = path.read_text(encoding="utf-8", errors="ignore")
                total += self.ingest_text(path.stem, content, "file", {"path": str(path), "collection": "documents"})
        return total
